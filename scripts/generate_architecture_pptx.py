from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT_PATH = Path("docs/self_improving_agent_architecture.pptx")
OUT_PATH_ROOT = Path("self_evolving_agents.pptx")

BG = RGBColor(245, 247, 250)
NAVY = RGBColor(14, 34, 64)
BLUE = RGBColor(48, 102, 190)
TEAL = RGBColor(34, 142, 147)
GOLD = RGBColor(215, 154, 53)
RED = RGBColor(184, 63, 69)
SLATE = RGBColor(76, 91, 106)
LIGHT = RGBColor(230, 236, 242)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(31, 41, 55)


def set_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12.1), Inches(0.9))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Aptos Display"
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = NAVY
    if subtitle:
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = subtitle
        r.font.name = "Aptos"
        r.font.size = Pt(12)
        r.font.color.rgb = SLATE


def add_footer(slide, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.0), Inches(12.0), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = text
    r.font.name = "Aptos"
    r.font.size = Pt(9)
    r.font.color.rgb = SLATE


def add_panel(slide, left: float, top: float, width: float, height: float, fill: RGBColor, title: str, body: list[str], title_color: RGBColor = NAVY) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill

    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Aptos"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = title_color

    for line in body:
        p = tf.add_paragraph()
        p.level = 0
        p.bullet = True
        r = p.add_run()
        r.text = line
        r.font.name = "Aptos"
        r.font.size = Pt(11)
        r.font.color.rgb = DARK


def add_box(slide, left: float, top: float, width: float, height: float, title: str, body: str, fill: RGBColor = WHITE, line: RGBColor = BLUE, title_fill: RGBColor | None = None) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.3)

    if title_fill is not None:
        header = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(left + 0.05),
            Inches(top + 0.05),
            Inches(width - 0.1),
            Inches(0.36),
        )
        header.fill.solid()
        header.fill.fore_color.rgb = title_fill
        header.line.color.rgb = title_fill
        htf = header.text_frame
        p = htf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.name = "Aptos"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = WHITE

        body_top = top + 0.47
        body_height = height - 0.52
        body_shape = slide.shapes.add_textbox(Inches(left + 0.14), Inches(body_top), Inches(width - 0.28), Inches(body_height))
        tf = body_shape.text_frame
        tf.word_wrap = True
        for i, line in enumerate(body.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = line
            r.font.name = "Aptos"
            r.font.size = Pt(10.5)
            r.font.color.rgb = DARK
        return

    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Aptos"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = NAVY

    for i, line in enumerate(body.split("\n")):
        p = tf.add_paragraph()
        p.bullet = i > -1
        r = p.add_run()
        r.text = line
        r.font.name = "Aptos"
        r.font.size = Pt(10.5)
        r.font.color.rgb = DARK


def add_arrow(slide, left: float, top: float, width: float, height: float, text: str = "") -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.CHEVRON,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = LIGHT
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = "Aptos"
    r.font.size = Pt(10)
    r.font.color.rgb = SLATE


def add_callout(slide, left: float, top: float, width: float, height: float, text: str, fill: RGBColor, font_color: RGBColor = WHITE) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = "Aptos"
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = font_color


def make_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def slide_title(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)

    banner = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2)
    )
    banner.fill.solid()
    banner.fill.fore_color.rgb = NAVY
    banner.line.color.rgb = NAVY

    box = slide.shapes.add_textbox(Inches(0.65), Inches(1.45), Inches(12.0), Inches(1.4))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Architecture of the Self-Improving Agent"
    r.font.name = "Aptos Display"
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = NAVY
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "How the reusable evoagent framework, the harness layer, and the evolving context work together"
    r.font.name = "Aptos"
    r.font.size = Pt(15)
    r.font.color.rgb = SLATE

    add_panel(
        slide, 0.7, 3.0, 3.85, 2.45, LIGHT, "Design Objective",
        [
            "Keep the model fixed; evolve the system around it.",
            "Reuse generic agent-improvement patterns via evoagent.",
            "Continuously learn from real user runs instead of one-off eval sweeps.",
        ],
    )
    add_panel(
        slide, 4.75, 3.0, 3.85, 2.45, RGBColor(225, 239, 255), "Three Learning Surfaces",
        [
            "Harness: middleware, tool parameters, completion checks.",
            "Context: prompts, skills, episodic memory, semantic memory.",
            "Outer loop: structural code changes when the inner loop plateaus.",
        ],
    )
    add_panel(
        slide, 8.8, 3.0, 3.85, 2.45, RGBColor(227, 243, 238), "Core Data Plane",
        [
            "LangSmith traces + local traces capture what happened.",
            "Graders produce numerical and textual feedback.",
            "Run log and version stores turn traces into durable learning signal.",
        ],
    )
    add_footer(slide, "Generated from current repository architecture")


def slide_influences(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "Research Lineage", "The implementation combines ideas from three reference directions plus LangSmith trace instrumentation.")

    add_box(slide, 0.55, 1.35, 4.0, 1.65, "Karpathy autoresearch", "Outer loop around a real task.\nAgent proposes improvements, evaluates them, keeps what works.\nThis project adapts that pattern into continuous agent improvement rather than GPU training.", fill=WHITE, line=BLUE, title_fill=BLUE)
    add_box(slide, 4.67, 1.35, 4.0, 1.65, "Meta-Harness", "Optimize the harness using execution traces, scores, and prior versions.\nThis project adopts trace-driven diagnosis and safe proposal selection, but does it continuously from live runs.", fill=WHITE, line=TEAL, title_fill=TEAL)
    add_box(slide, 8.79, 1.35, 4.0, 1.65, "Continual learning for agents", "Separate model, harness, and context as distinct learning surfaces.\nThis project mainly learns in token space: prompts, skills, and memory.", fill=WHITE, line=GOLD, title_fill=GOLD)

    add_box(slide, 0.85, 3.45, 12.0, 2.45, "Project-specific synthesis", "1. `evoagent/` packages the reusable framework: protocols, memory, skills, middleware, grading, tracing, and offline evolution helpers.\n2. `src/` wires that framework into a research agent built on `deepagents`.\n3. `src/evolution/daemon.py` is the always-on optimizer that consumes real run results.\n4. LangSmith traces enrich both prompt optimization and harness optimization with higher-fidelity execution evidence.", fill=RGBColor(252, 252, 252), line=SLATE, title_fill=NAVY)
    add_callout(slide, 4.55, 6.15, 4.2, 0.45, "Key idea: traces are the common currency", NAVY)
    add_footer(slide, "Refs: Meta-Harness, LangChain continual learning, karpathy/autoresearch, LangSmith traces")


def slide_layers(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "Reusable evoagent Design", "The framework is intentionally layered so other agent types can reuse only the pieces they need.")

    add_box(slide, 0.65, 1.45, 12.0, 0.75, "Layering principle", "Layer 0: protocols and types. Layer 1: memory and skills. Layer 2: harness, graders, tracing. Layer 3: evolution utilities. The application code in `src/` sits above this and supplies the domain-specific agent factory, prompts, tools, and CLI.", fill=WHITE, line=SLATE, title_fill=SLATE)
    add_box(slide, 0.65, 2.45, 2.9, 3.25, "Layer 0\n`evoagent/core/`", "Shared types such as `GraderResult`, `TaskResult`, `TrajectoryMetrics`.\nProtocols such as `Grader`, `AgentFactory`, `MemoryBackend`, `SkillStore`, `PromptStore`, `MiddlewareHook`.\nThis is what makes the package composable.", fill=RGBColor(233, 242, 255), line=BLUE)
    add_box(slide, 3.78, 2.45, 2.9, 3.25, "Layer 1\nMemory + Skills", "File-backed memory namespaces.\nContext compression and semantic dedup.\nSKILL.md discovery, creation, validation, and extraction from runs.", fill=RGBColor(226, 244, 241), line=TEAL)
    add_box(slide, 6.91, 2.45, 2.9, 3.25, "Layer 2\nHarness + Tracing + Grading", "Middleware for self-verification, context injection, loop detection, time budget, reasoning phases, trace capture.\nTrace capture/fetching.\nEvaluation surfaces for quality, efficiency, task completion, and verification.", fill=RGBColor(255, 244, 226), line=GOLD)
    add_box(slide, 10.04, 2.45, 2.6, 3.25, "Layer 3\nEvolution", "Prompt optimization.\nError analysis.\nSleep review.\nBatch evolution helpers for propose-run-grade-accept loops.", fill=RGBColor(248, 230, 231), line=RED)
    add_footer(slide, "Main framework boundary: everything under `evoagent/` is reusable")


def slide_runtime(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "Application Runtime Architecture", "The `src/` package specializes the framework into a self-improving research agent.")

    add_box(slide, 0.55, 1.55, 2.25, 1.2, "Version Stores", "PromptStore returns LATEST promoted prompt (ratchet).\nHarnessConfigStore returns LATEST promoted config.\nFrozen `promotion_score` per version.", fill=WHITE, line=BLUE, title_fill=BLUE)
    add_arrow(slide, 2.95, 1.92, 0.65, 0.42)
    add_box(slide, 3.75, 1.35, 2.6, 1.6, "Agent Factory\n`src/agent/deep_agent.py`", "Creates LLM.\nCreates search tool with harness-tuned parameters.\nLoads prompt, memory context, skills, subagents, middleware.", fill=WHITE, line=TEAL, title_fill=TEAL)
    add_arrow(slide, 6.5, 1.92, 0.65, 0.42)
    add_box(slide, 7.25, 1.35, 2.65, 1.6, "deepagents Runtime", "System prompt = current prompt + compressed memory context.\nMiddleware stack from `default_middleware_stack()`.\nPrimary tool = search.", fill=WHITE, line=GOLD, title_fill=GOLD)
    add_arrow(slide, 10.03, 1.92, 0.65, 0.42)
    add_box(slide, 10.8, 1.55, 1.95, 1.2, "Task Output", "Final answer\nplus trace side effects", fill=WHITE, line=RED, title_fill=RED)

    add_box(slide, 0.7, 3.35, 3.85, 2.5, "What is fixed vs learned", "Fixed at runtime:\n`deepagents` execution model, grader classes, storage schema.\n\nLearned across runs:\nprompt versions, harness config versions, learned skills, episodic summaries, semantic patterns, meta-instructions.", fill=RGBColor(252, 252, 252), line=SLATE)
    add_box(slide, 4.8, 3.35, 3.85, 2.5, "Why `src/` stays thin", "Application code mostly composes library pieces rather than re-implementing them.\nThat is what makes the framework reusable: a different agent can swap the toolset and prompt domain while keeping the same improvement machinery.", fill=RGBColor(252, 252, 252), line=SLATE)
    add_box(slide, 8.9, 3.35, 3.75, 2.5, "Critical join points", "`create_agent()` joins prompt store, harness config store, skills directory, memory compression, and tracing.\nThat function is the bridge where offline learning re-enters the hot path of the next run.", fill=RGBColor(252, 252, 252), line=SLATE)
    add_footer(slide, "Hot path: read best versions -> run agent -> emit new evidence")


def slide_middleware(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "Harness Stack", "`evoagent/harness/builder.py` assembles six middleware modules; `HarnessConfig` controls most of their behavior.")

    x_positions = [0.45, 2.55, 4.65, 6.75, 8.85, 10.95]
    titles = [
        "1. SelfVerification",
        "2. ContextAssembly",
        "3. LoopDetection",
        "4. TimeBudget",
        "5. ReasoningSandwich",
        "6. TraceCapture",
    ]
    bodies = [
        "Checks structure, length, task coverage, and optional completion checks; retries before finalizing.",
        "Injects environment snapshot and discovered skills on the first model call.",
        "Detects repeated searches, repeated tools, and file-edit loops; nudges toward synthesis.",
        "Appends escalating warnings to tool responses as budget is consumed.",
        "Allocates more reasoning effort to planning and verification phases than execution phase.",
        "Persists trace data for offline diagnosis and LangSmith/local analysis.",
    ]
    fills = [RGBColor(233, 242, 255), RGBColor(226, 244, 241), RGBColor(255, 244, 226), RGBColor(248, 230, 231), RGBColor(232, 236, 252), RGBColor(236, 236, 236)]
    for i, x in enumerate(x_positions):
        add_box(slide, x, 1.65, 1.85, 3.7, titles[i], bodies[i], fill=fills[i], line=SLATE)
        if i < len(x_positions) - 1:
            add_arrow(slide, x + 1.92, 3.05, 0.15, 0.25)

    add_box(slide, 0.8, 5.75, 12.0, 0.9, "Tunable surface", "`HarnessConfig` includes retries, required sections, completion checks, loop thresholds, file edit thresholds, warning points, reasoning efforts, planning call count, environment detection, and search-tool parameters such as retry count, delay, result count, and search depth.", fill=WHITE, line=NAVY, title_fill=NAVY)
    add_footer(slide, "This harness is learned configuration, not hard-coded policy")


def slide_single_run(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "How One Run Produces Learning Signal", "The run path is implemented in `src/cli/commands.py::cmd_run`.")

    add_box(slide, 0.55, 1.6, 2.0, 1.0, "1. Execute", "Agent answers the task with current prompt, memory context, tools, and harness.", fill=WHITE, line=BLUE, title_fill=BLUE)
    add_arrow(slide, 2.65, 1.95, 0.55, 0.3)
    add_box(slide, 3.3, 1.6, 2.0, 1.0, "2. Grade", "Trajectory is scored for task completion, quality, efficiency, and claim verification.", fill=WHITE, line=TEAL, title_fill=TEAL)
    add_arrow(slide, 5.4, 1.95, 0.55, 0.3)
    add_box(slide, 6.05, 1.6, 2.0, 1.0, "3. Observed score", "Per-run scores feed the running observed average.\nThe frozen `promotion_score` is NOT touched here.", fill=WHITE, line=GOLD, title_fill=GOLD)
    add_arrow(slide, 8.15, 1.95, 0.55, 0.3)
    add_box(slide, 8.8, 1.6, 2.0, 1.0, "4. Run log", "Structured entry is appended to `evolution_state/run_log.jsonl`.", fill=WHITE, line=RED, title_fill=RED)
    add_arrow(slide, 10.9, 1.95, 0.55, 0.3)
    add_box(slide, 11.55, 1.6, 1.25, 1.0, "5. Reflect", "Store episodic + semantic memory.", fill=WHITE, line=SLATE, title_fill=SLATE)

    add_box(slide, 0.75, 3.1, 5.9, 2.8, "What gets persisted immediately", "`update_score()` only mutates the running observed average. The `promotion_score` field is frozen at promotion time, so observed drift can never silently roll the active version backward (the ratchet).\n\n`RunLogEntry` captures task, output, classification, average score, per-grader reasoning, prompt version, harness config version, and optional trace path.\n\n`reflect_and_store()` writes:\n- episodic memory = run summary, what worked, what to improve, output preview\n- semantic memory = extracted facts and patterns", fill=RGBColor(252, 252, 252), line=SLATE)
    add_box(slide, 6.95, 3.1, 5.65, 2.8, "Why this matters", "The project does not wait for a separate benchmark stage to learn.\nEvery user task is simultaneously:\n- serving traffic\n- producing reward signal\n- producing text feedback\n- producing memory\n- producing trace evidence for later diagnosis", fill=RGBColor(252, 252, 252), line=SLATE)
    add_footer(slide, "Per-run loop = hot-path learning and evidence capture")


def slide_daemon(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "Background Daemon", "`src/evolution/daemon.py` is the primary continuous-improvement loop.")

    add_box(slide, 0.65, 1.55, 2.15, 1.1, "Trigger", "Poll `run_log.jsonl` every 60s.\nStart when 3+ unprocessed runs accumulate.\nLoad held-out tasks via `_load_holdout_tasks()` (stable seed=42).", fill=WHITE, line=BLUE, title_fill=BLUE)
    add_arrow(slide, 2.95, 1.92, 0.55, 0.3)
    add_box(slide, 3.55, 1.3, 2.0, 1.6, "1. Prompt opt", "Analyze failures + traces.\nGenerate candidates.\nGate via held-out pairwise.\nNo orphan writes on reject.", fill=WHITE, line=TEAL, title_fill=TEAL)
    add_arrow(slide, 5.7, 1.92, 0.45, 0.3)
    add_box(slide, 6.2, 1.3, 1.75, 1.6, "2. Success skills", "Extract reusable patterns from high-scoring runs.", fill=WHITE, line=GOLD, title_fill=GOLD)
    add_arrow(slide, 8.08, 1.92, 0.45, 0.3)
    add_box(slide, 8.58, 1.3, 1.75, 1.6, "3. Failure skills", "Create defensive skills from bad runs.", fill=WHITE, line=RED, title_fill=RED)
    add_arrow(slide, 10.46, 1.92, 0.45, 0.3)
    add_box(slide, 10.96, 1.3, 1.75, 1.6, "4. Memory compression", "Deduplicate semantic memory.", fill=WHITE, line=SLATE, title_fill=SLATE)

    add_box(slide, 1.15, 3.55, 11.0, 2.15, "5. Harness optimization (also gated)", "After prompt/skill/memory updates, `optimize_harness()` reads grading summaries + trace evidence, proposes changes to `HarnessConfig`, then runs `validate_candidate_harness()` against the SAME held-out task set: each task is run once with the current config and once with the candidate, with the same prompt, then pairwise compared. Only proposals that win strict majority on the holdout are saved.\n\nWhen the gate rejects, the optimizer returns the unchanged version number — no file is written and no phantom 'Harness upgraded' line appears in the daemon log.", fill=RGBColor(252, 252, 252), line=NAVY, title_fill=NAVY)
    add_footer(slide, "Cold-start path: `src evolve` bootstraps; the daemon handles all live improvement after that")


def slide_harness_update(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "How the Harness Gets Updated", "This is the most explicit embodiment of the Meta-Harness idea in the project.")

    add_box(slide, 0.65, 1.5, 2.9, 1.65, "Diagnostic inputs", "Run-log entries provide scores, classifications, and per-grader reasons.\nTrace fetcher prefers LangSmith traces, falls back to local JSON.\nOptimizer sees both metrics and execution details.", fill=WHITE, line=BLUE, title_fill=BLUE)
    add_box(slide, 3.85, 1.5, 2.9, 1.65, "LLM proposal phase", "Optimizer prompt includes current config, per-dimension failure rates, top failure reasons, trace data.\nLLM emits JSON changes, not free-form code edits.\nMost-conservative-pick across multiple candidates.", fill=WHITE, line=TEAL, title_fill=TEAL)
    add_box(slide, 7.05, 1.5, 2.6, 1.65, "Safety layer", "Unknown parameters ignored.\nEffort strings validated.\nNumeric values clamped to safe bounds.\nInvalid candidates dropped.", fill=WHITE, line=GOLD, title_fill=GOLD)
    add_box(slide, 9.95, 1.5, 2.75, 1.65, "Held-out promotion gate", "`validate_candidate_harness()` runs each held-out task with BOTH old and new config (same prompt). Pairwise judge compares outputs. Strict majority required.\nLoser is not saved.", fill=WHITE, line=RED, title_fill=RED)

    add_box(slide, 0.95, 3.7, 11.8, 2.15, "What can evolve", "Middleware behavior:\n`max_retries`, `min_length`, `verify_against_task`, `required_sections`, `completion_checks`, `max_similar`, `max_total`, `max_file_edits`, `max_repeated_tools`, `budget_seconds`, `warn_at`, `planning_effort`, `implementation_effort`, `verification_effort`, `planning_calls`, `detect_env`.\n\nTool behavior:\n`search_max_retries`, `search_retry_delay`, `search_max_results`, `search_depth`.\n\nBecause the search tool is created with `harness_cfg`, tool-level behavior and harness-level behavior evolve together.", fill=RGBColor(252, 252, 252), line=SLATE)
    add_footer(slide, "Updated harness versions are JSON config files, so rollback and inspection stay simple")


def slide_context(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "How Context Evolves Dynamically", "The project learns in token space: prompts, skills, memories, and meta-instructions are treated as mutable context.")

    add_box(slide, 0.55, 1.55, 3.05, 1.9, "Prompt evolution", "Prompt versions live in `prompts/v*.json`.\nEach run updates the version's running observed score.\nDaemon adds a new version only after held-out pairwise validation passes.\n`get_current_prompt()` returns the LATEST promoted version (ratchet).", fill=WHITE, line=BLUE, title_fill=BLUE)
    add_box(slide, 3.85, 1.55, 3.05, 1.9, "Skill evolution", "Success runs create reusable skills.\nFailure runs create defensive skills.\n`ContextAssemblyMiddleware` discovers skills and injects the available skill catalog into the first model call.", fill=WHITE, line=TEAL, title_fill=TEAL)
    add_box(slide, 7.15, 1.55, 2.75, 1.9, "Memory evolution", "Reflection stores episodic summaries and semantic facts/patterns.\nCompression prioritizes meta-instructions, then learned patterns, then recent run summaries under a token budget.", fill=WHITE, line=GOLD, title_fill=GOLD)
    add_box(slide, 10.15, 1.55, 2.55, 1.9, "Sleep-time review", "Cross-run trace analysis writes `meta_instruction` memories and can remove contradictory semantic memories.", fill=WHITE, line=RED, title_fill=RED)

    add_box(slide, 0.9, 4.0, 11.7, 1.85, "Injection path back into the next run", "1. `create_agent()` calls `compress_context(memory_store, task, token_budget)`.\n2. The compressed memory string is inserted into the prompt template via `{memory_context}`.\n3. `ContextAssemblyMiddleware.before_model()` adds environment context and available skill descriptions on the first model call.\n4. The result is a run-time prompt that already contains the system's latest learned strategy, facts, and reusable tactics before the model makes its first major decision.", fill=RGBColor(252, 252, 252), line=NAVY, title_fill=NAVY)
    add_footer(slide, "The model weights stay fixed; the context around the model changes continuously")


def slide_langsmith(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "Role of LangSmith Traces", "LangSmith is used as the high-fidelity observation layer for improving the agent.")

    add_box(slide, 0.65, 1.5, 3.8, 1.95, "TraceFetcher behavior", "Lazy-initializes a LangSmith client from project settings.\nCan fetch recent root runs, full descendant traces for a run, traces matching a task prefix, or batched traces for run-log entries.\nFalls back to local trace JSON when unavailable.", fill=WHITE, line=BLUE, title_fill=BLUE)
    add_box(slide, 4.75, 1.5, 3.8, 1.95, "Prompt optimizer use", "Failed and partial runs are paired with rich traces.\nTrace digests surface search queries, errors, tool-call decisions, and timing behavior.\nThat gives the metaprompt more specific failure diagnoses than score summaries alone.", fill=WHITE, line=TEAL, title_fill=TEAL)
    add_box(slide, 8.85, 1.5, 3.8, 1.95, "Harness optimizer use", "The harness optimizer sees the current config, per-dimension failure stats, and trace evidence.\nThis is how tool retry counts, loop thresholds, reasoning allocation, and completion checks get tuned based on actual behavior.", fill=WHITE, line=GOLD, title_fill=GOLD)

    add_box(slide, 1.05, 4.0, 11.2, 1.75, "Why traces matter architecturally", "Scores tell the system that something was wrong; traces tell it what actually happened.\nThe project uses LangSmith traces as the bridge from opaque outcome metrics to actionable prompt and harness updates. That is the main reason the system can do continuous improvement without retraining the model.", fill=RGBColor(252, 252, 252), line=SLATE)
    add_footer(slide, "If LangSmith is absent, the system degrades gracefully to local trace files")


def slide_closed_loop(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "Closed-Loop Dynamic Evolution", "The architecture is best understood as a multi-speed feedback system.")

    add_box(slide, 0.55, 1.55, 3.85, 3.95, "Fast loop: per task", "Input task.\nAgent runs with best prompt + best harness config + current skills + compressed memory.\nMiddleware shapes execution in real time.\nGrading, scoring, reflection, trace capture, and run-log append happen immediately.", fill=RGBColor(233, 242, 255), line=BLUE, title_fill=BLUE)
    add_box(slide, 4.75, 1.55, 3.85, 3.95, "Medium loop: background daemon", "When enough fresh runs accumulate, the daemon turns evidence into updated prompt versions, new skills, compressed memory, and new harness configs.\nThose artifacts become part of the next serving run automatically.", fill=RGBColor(226, 244, 241), line=TEAL, title_fill=TEAL)
    add_box(slide, 8.95, 1.55, 3.85, 3.95, "Slow loop: outer loop", "When inner-loop improvements plateau, the codebase stores evolution state so a coding agent can make structural changes: new tools, different models, refactors, architecture changes.\nThis mirrors Karpathy-style code evolution.", fill=RGBColor(255, 244, 226), line=GOLD, title_fill=GOLD)

    add_callout(slide, 2.0, 5.95, 9.3, 0.5, "Net effect: the system changes the instructions, controls, and remembered knowledge around the model until the next task is solved under a better operating regime.", NAVY)
    add_footer(slide, "This is continuous systems learning, not one-shot prompt tuning")


def slide_ratchet(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "Promotion Ratchet", "Active versions move only forward; observed-score drift cannot silently roll the champion back.")

    add_box(slide, 0.55, 1.55, 4.0, 2.4, "Frozen `promotion_score`", "Set once at promotion time, never mutated.\nNew champions inherit the parent's `promotion_score` so each promotion has a stable floor.\nUsed by the optimizer for monitoring and as the baseline for the next gate.", fill=WHITE, line=BLUE, title_fill=BLUE)
    add_box(slide, 4.7, 1.55, 4.0, 2.4, "Running `score`", "Incremental average of post-promotion run scores via `update_score()`.\nDrifts up or down with new evidence.\nServes as the optimizer's signal to ATTEMPT improvement, but does not decide which version is active.", fill=WHITE, line=TEAL, title_fill=TEAL)
    add_box(slide, 8.85, 1.55, 4.0, 2.4, "Active = LATEST", "`get_current_prompt()` and `load_best()` return the latest version unconditionally.\nNot max-scored.\nA promotion only happens through the held-out gate, so 'latest' is by construction the most recently validated champion.", fill=WHITE, line=GOLD, title_fill=GOLD)

    add_box(slide, 0.85, 4.15, 12.0, 1.7, "What this prevents", "Before the ratchet, a string of unlucky or buggy runs (e.g. the recent `extract_output` regression where phantom self-check strings were graded as the agent's output) would drag the champion's running score below an older version's, and `max(score)` selection would silently activate the older version. The active prompt could change without any optimizer involvement.\n\nWith the ratchet, observed scores can swing freely but the active version only moves forward through an explicit, gated promotion. The system self-corrects through the front door (more optimization) instead of through the back door (silent rollback).", fill=RGBColor(252, 252, 252), line=NAVY, title_fill=NAVY)
    add_callout(slide, 3.5, 6.15, 6.3, 0.45, "Karpathy ratchet: separate observation from commitment", NAVY)
    add_footer(slide, "Implementation: `src/agent/prompt_store.py`, `src/evolution/harness_config.py`")


def slide_gates(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "Held-out Promotion Gates", "Both prompt and harness optimizers must beat the champion on a stable held-out task set.")

    add_box(slide, 0.55, 1.5, 4.0, 2.5, "Held-out source", "`tasks/research_tasks.json` is split via `_split_tasks` (random.Random(42), HOLDOUT_FRACTION=0.3).\nThe holdout slice is stable across cycles, so trends are comparable and the optimizer never validates on its own training input.\n`_load_holdout_tasks()` reads it once per cycle in the daemon.", fill=WHITE, line=BLUE, title_fill=BLUE)
    add_box(slide, 4.7, 1.5, 4.0, 2.5, "Apples-to-apples pairwise", "For each held-out task: run the CURRENT champion and the CANDIDATE fresh, in the same cycle, with the same harness state (or same prompt for the harness gate).\nA pairwise LLM judge picks the better output. Position is randomized to defeat slot bias.\nStrict majority required: `wins_new > wins_old`.", fill=WHITE, line=TEAL, title_fill=TEAL)
    add_box(slide, 8.85, 1.5, 4.0, 2.5, "Defensive errors + log tags", "Run errors on either side count as a loss for the candidate — the system never promotes a config or prompt that crashes the agent.\nLog tags `[holdout]` (prompt) and `[harness-holdout]` (harness) make the active gate visible at a glance; `[batch-fallback]` flags the legacy overfit-prone path.", fill=WHITE, line=GOLD, title_fill=GOLD)

    add_box(slide, 0.85, 4.2, 12.0, 1.7, "What the gate catches", "The recent harness v2 incident is the canonical example. The optimizer hallucinated that 'broad/low-confidence research outputs' needed `min_length` loosened — but the failures were actually phantom SELF-CHECK strings from a separate `extract_output` bug. Without a gate, the optimizer's reasoning was self-validating: it saw bad outputs, proposed a fix, and the fix was promoted blindly. With the held-out gate, that same proposal would have to actually win pairwise on the holdout. It would not, because the change does not address the real failure mode.\n\nThe gate makes the optimizer pay the cost of being right.", fill=RGBColor(252, 252, 252), line=NAVY, title_fill=NAVY)
    add_callout(slide, 3.5, 6.2, 6.3, 0.45, "Promotion is now genuinely monotonic", NAVY)
    add_footer(slide, "Implementation: `src/evolution/prompt_optimizer.py::_validate_on_holdout`, `harness_optimizer.py::validate_candidate_harness`")


def slide_files(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "Key Files to Read", "These are the implementation anchors behind the architecture.")

    add_box(slide, 0.65, 1.45, 4.0, 4.8, "Reusable framework (`evoagent/`)", "`evoagent/core/protocols.py`\nContracts for pluggability.\n\n`evoagent/harness/builder.py`\nDefault middleware assembly.\n\n`evoagent/harness/middleware.py`\nActual runtime controls.\n\n`evoagent/memory/store.py`\nPersistent episodic and semantic memory.\n\n`evoagent/memory/compression.py`\nContext assembly and dedup.\n\n`evoagent/skills/extractor.py`\nSuccess and failure skill learning.\n\n`evoagent/evolution/sleep_review.py`\nOffline meta-instruction learning.", fill=WHITE, line=BLUE, title_fill=BLUE)
    add_box(slide, 4.9, 1.45, 4.0, 4.8, "Application wiring (`src/`)", "`src/agent/deep_agent.py`\nBuilds the actual research agent from learned artifacts.\n\n`src/cli/commands.py`\nSingle-run execution, grading, reflection, and run logging.\n\n`src/evolution/daemon.py`\nContinuous optimizer.\n\n`src/evolution/harness_optimizer.py`\nSafe config proposals from traces.\n\n`src/evolution/prompt_optimizer.py`\nPrompt revision from failure analysis.\n\n`src/tracing/fetcher.py`\nLangSmith integration.", fill=WHITE, line=TEAL, title_fill=TEAL)
    add_box(slide, 9.15, 1.45, 3.45, 4.8, "State directories", "`prompts/`\nVersioned prompt JSON.\n\n`harness_config/`\nVersioned harness JSON.\n\n`memory/episodic` and `memory/semantic`\nLearned context.\n\n`skills/`\nGenerated SKILL.md patterns.\n\n`evolution_state/run_log.jsonl`\nQueue of evidence for the daemon.\n\n`traces/`\nLocal fallback trace store.", fill=WHITE, line=GOLD, title_fill=GOLD)
    add_footer(slide, "These directories are the durable memory of the system")


def slide_refs(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_title(slide, "References", "Primary design influences and repository documentation used for this deck.")

    refs = [
        "Repository docs: `architecture.md`, `README.md`, `docs/superpowers/specs/2026-04-05-evoagent-library-design.md`, `docs/superpowers/specs/2026-04-05-harness-improvements-design.md`",
        "Meta-Harness project page: https://yoonholee.com/meta-harness/",
        "LangChain blog, 'Continual learning for AI agents' (Apr 5, 2026): https://blog.langchain.com/continual-learning-for-ai-agents/",
        "karpathy/autoresearch repository: https://github.com/karpathy/autoresearch",
        "LangSmith traces are integrated via `src/tracing/fetcher.py` and are referenced throughout the optimization loop.",
    ]

    add_box(slide, 0.8, 1.6, 11.8, 4.6, "Reference list", "\n".join(refs), fill=WHITE, line=NAVY, title_fill=NAVY)
    add_callout(slide, 3.25, 6.45, 6.75, 0.45, "Deck generator: `scripts/generate_architecture_pptx.py`", SLATE)
    add_footer(slide, "Generated PowerPoint output: docs/self_improving_agent_architecture.pptx")


def main() -> None:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs = make_prs()
    slide_title(prs)
    slide_influences(prs)
    slide_layers(prs)
    slide_runtime(prs)
    slide_middleware(prs)
    slide_single_run(prs)
    slide_daemon(prs)
    slide_harness_update(prs)
    slide_context(prs)
    slide_langsmith(prs)
    slide_closed_loop(prs)
    slide_ratchet(prs)
    slide_gates(prs)
    slide_files(prs)
    slide_refs(prs)
    prs.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")
    # Also write to the project-root copy that is referenced when sharing
    # the deck outside the docs folder. Both files stay in sync.
    prs.save(OUT_PATH_ROOT)
    print(f"Wrote {OUT_PATH_ROOT}")


if __name__ == "__main__":
    main()
